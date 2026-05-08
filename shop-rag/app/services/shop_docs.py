from __future__ import annotations

import hashlib
import html
import json
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import httpx
from pydantic import BaseModel, Field

from app.core.config import Settings
from app.qdrant.client import create_qdrant_client
from app.qdrant.importer import (
    DEFAULT_IMPORT_EMBEDDING_MODEL,
    ChunkRecord,
    QdrantChunkImporter,
    SentenceTransformerEmbeddingProvider,
)
from app.services.job_manager import JobContext
from app.services.source_index import SourceIndex, SourceRecord


SHOP_DOCS_COLLECTION = "shop_docs_v1"
SHOP_DOCS_STAGING_COLLECTION = "shop_docs_v1_staging"
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".html", ".htm"}


class ShopDocScanItem(BaseModel):
    source_ref: str
    source_type: str
    supported: bool = True
    warning: str = ""


class ShopDocScanResult(BaseModel):
    inbox_path: str
    supported_count: int
    unsupported_count: int
    items: list[ShopDocScanItem]


class ExtractedPage(BaseModel):
    page_number: int
    text: str


class ExtractedDocument(BaseModel):
    title: str
    source_ref: str
    source_url: str | None = None
    file_type: str
    text: str
    content_hash: str
    pages: list[ExtractedPage] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)


@dataclass
class ShopDocsService:
    inbox_path: Path

    def scan(self) -> ShopDocScanResult:
        self.inbox_path.mkdir(parents=True, exist_ok=True)
        items: list[ShopDocScanItem] = []
        for path in sorted(self.inbox_path.rglob("*")):
            if path.is_dir():
                continue
            relative_ref = path.relative_to(self.inbox_path).as_posix()
            if path.name.lower().endswith(".metadata.json"):
                continue
            if path.name.lower() == "urls.txt":
                items.extend(self._scan_urls(path))
                continue
            suffix = path.suffix.lower()
            if suffix in SUPPORTED_EXTENSIONS:
                items.append(ShopDocScanItem(source_ref=relative_ref, source_type=suffix.lstrip(".")))
        return ShopDocScanResult(
            inbox_path=str(self.inbox_path),
            supported_count=sum(1 for item in items if item.supported),
            unsupported_count=sum(1 for item in items if not item.supported),
            items=items,
        )

    def extract_local_file(self, path: Path) -> ExtractedDocument:
        metadata = _metadata_for_file(path)
        suffix = path.suffix.lower()
        pages: list[ExtractedPage] = []
        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif suffix in {".html", ".htm"}:
            text = _html_to_text(path.read_text(encoding="utf-8", errors="ignore"))
        elif suffix == ".pdf":
            pages = _extract_pdf_pages(path)
            text = "\n\n".join(page.text for page in pages)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".xlsx":
            text = _extract_xlsx(path)
        else:
            raise ValueError(f"Unsupported file type: {path.suffix}")
        cleaned = _clean_text(text)
        title = str(metadata.get("title") or (_title_from_html(path.read_text(encoding="utf-8", errors="ignore")) if suffix in {".html", ".htm"} else "") or path.stem)
        return ExtractedDocument(
            title=title,
            source_ref=path.relative_to(self.inbox_path).as_posix() if path.is_relative_to(self.inbox_path) else path.name,
            source_url=str(metadata.get("source_url") or "") or None,
            file_type=suffix.lstrip("."),
            text=cleaned,
            content_hash=_hash_text(cleaned),
            pages=_clean_pages(pages) if suffix == ".pdf" else [],
            metadata={"path": str(path), **metadata},
        )

    def extract_url(self, url: str) -> ExtractedDocument:
        response = httpx.get(url, timeout=30, follow_redirects=True)
        response.raise_for_status()
        content_type = response.headers.get("content-type", "").lower()
        if "application/pdf" in content_type or url.lower().split("?")[0].endswith(".pdf"):
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
                handle.write(response.content)
                temp_path = Path(handle.name)
            try:
                pages = _extract_pdf_pages(temp_path)
            finally:
                temp_path.unlink(missing_ok=True)
            text = "\n\n".join(page.text for page in pages)
            cleaned = _clean_text(text)
            return ExtractedDocument(
                title=Path(url.split("?")[0]).stem or url,
                source_ref=url,
                source_url=url,
                file_type="pdf",
                text=cleaned,
                content_hash=_hash_text(cleaned),
                pages=_clean_pages(pages),
                metadata={"url": url, "content_type": content_type},
            )
        text = _html_to_text(response.text)
        title = _title_from_html(response.text) or url
        cleaned = _clean_text(text)
        return ExtractedDocument(
            title=title,
            source_ref=url,
            source_url=url,
            file_type="url",
            text=cleaned,
            content_hash=_hash_text(cleaned),
            metadata={"url": url},
        )

    def extract_all(self, context: JobContext | None = None) -> list[ExtractedDocument]:
        documents: list[ExtractedDocument] = []
        for item in self.scan().items:
            if not item.supported:
                continue
            if item.source_type == "url":
                if context:
                    context.log(f"Fetching URL: {item.source_ref}")
                documents.append(self.extract_url(item.source_ref))
            else:
                path = self.inbox_path / item.source_ref
                if context:
                    context.log(f"Extracting file: {path.name}")
                documents.append(self.extract_local_file(path))
        return [document for document in documents if document.text.strip()]

    def _scan_urls(self, path: Path) -> list[ShopDocScanItem]:
        items: list[ShopDocScanItem] = []
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            url = line.strip()
            if not url or url.startswith("#"):
                continue
            supported = url.lower().startswith(("http://", "https://"))
            items.append(
                ShopDocScanItem(
                    source_ref=url,
                    source_type="url",
                    supported=supported,
                    warning="" if supported else "URL must start with http:// or https://",
                )
            )
        return items


def run_shop_docs_import(settings: Settings, source_index: SourceIndex, context: JobContext) -> dict[str, object]:
    service = ShopDocsService(settings.shop_docs_inbox_path)
    context.set_step("extracting shop documents")
    documents = service.extract_all(context)
    if not documents:
        raise RuntimeError(f"No supported shop documents found in {settings.shop_docs_inbox_path}")

    context.set_step("chunking shop documents")
    chunks = _chunk_documents(documents)
    context.log(f"Created {len(chunks)} chunks from {len(documents)} documents")

    context.set_step("importing staging shop-doc collection")
    client = create_qdrant_client(settings)
    result = QdrantChunkImporter(
        client=client,
        embedding_provider=SentenceTransformerEmbeddingProvider(
            settings.query_embedding_model or DEFAULT_IMPORT_EMBEDDING_MODEL
        ),
        collection_name=SHOP_DOCS_STAGING_COLLECTION,
        batch_size=64,
        recreate=True,
    ).import_records(chunks)

    context.set_step("promoting shop-doc collection")
    QdrantChunkImporter(
        client=client,
        embedding_provider=SentenceTransformerEmbeddingProvider(
            settings.query_embedding_model or DEFAULT_IMPORT_EMBEDDING_MODEL
        ),
        collection_name=SHOP_DOCS_COLLECTION,
        batch_size=64,
        recreate=True,
    ).import_records(chunks)

    context.set_step("updating source index")
    source_index.replace_collection_sources(
        SHOP_DOCS_COLLECTION,
        [_source_record_from_doc(document, SHOP_DOCS_COLLECTION, chunks) for document in documents],
    )
    return {
        "collection": SHOP_DOCS_COLLECTION,
        "documents": len(documents),
        "chunks": len(chunks),
        "imported_points": result.imported_points,
    }


def _chunk_documents(documents: list[ExtractedDocument]) -> list[ChunkRecord]:
    chunks: list[ChunkRecord] = []
    for document in documents:
        if document.file_type == "pdf" and document.pages:
            chunks.extend(_chunk_pdf_document(document))
            continue
        parts = _split_text(document.text)
        for index, part in enumerate(parts):
            chunks.append(
                ChunkRecord(
                    chunk_id=f"{_hash_text(document.source_ref)}_chunk_{index}",
                    text=part,
                    metadata={
                        "title": document.title,
                        "category": str(document.metadata.get("category") or "Shop Docs"),
                        "source_url": document.source_url or "",
                        "content_url": document.source_ref,
                        "description": f"{document.metadata.get('category') or 'Shop document'}: {document.source_ref}",
                        "date_modified": "",
                        "chunk_index": index,
                        "total_chunks": len(parts),
                        "file_type": document.file_type,
                        "knowledge_type": "shop_docs",
                        "source_ref": document.source_ref,
                        "content_hash": document.content_hash,
                    },
                )
            )
    return chunks


def _chunk_pdf_document(document: ExtractedDocument) -> list[ChunkRecord]:
    chunks: list[ChunkRecord] = []
    chunk_index = 0
    page_parts: list[tuple[ExtractedPage, str]] = []
    for page in document.pages:
        for part in _split_text(page.text):
            page_parts.append((page, part))
    total_chunks = len(page_parts)
    for page, part in page_parts:
        source_url_with_page = _pdf_page_url(document.source_url, page.page_number)
        chunks.append(
            ChunkRecord(
                chunk_id=f"{_hash_text(document.source_ref)}_page_{page.page_number}_chunk_{chunk_index}",
                text=part,
                metadata={
                    "title": document.title,
                    "category": str(document.metadata.get("category") or "Shop Docs"),
                    "source_url": document.source_url or "",
                    "source_url_with_page": source_url_with_page or "",
                    "content_url": document.source_ref,
                    "description": f"{document.metadata.get('category') or 'Shop document'}: {document.source_ref}",
                    "date_modified": "",
                    "chunk_index": chunk_index,
                    "total_chunks": total_chunks,
                    "file_type": document.file_type,
                    "knowledge_type": "shop_docs",
                    "source_ref": document.source_ref,
                    "content_hash": document.content_hash,
                    "page_number": page.page_number,
                    "page_start": page.page_number,
                    "page_end": page.page_number,
                    "page_range": str(page.page_number),
                },
            )
        )
        chunk_index += 1
    return chunks


def _source_record_from_doc(document: ExtractedDocument, collection: str, chunks: list[ChunkRecord]) -> SourceRecord:
    chunk_count = sum(1 for chunk in chunks if chunk.metadata.get("source_ref") == document.source_ref)
    return SourceRecord(
        source_id=_hash_text(f"shop_docs:{document.source_ref}"),
        collection=collection,
        knowledge_type="shop_docs",
        title=document.title,
        source_ref=document.source_ref,
        source_url=document.source_url,
        file_type=document.file_type,
        category=str(document.metadata.get("category") or "Shop Docs"),
        chunk_count=chunk_count,
        content_hash=document.content_hash,
        metadata=document.metadata,
    )


def _split_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list[str]:
    text = text.strip()
    if len(text) <= chunk_size:
        return [text]
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        chunks.append(text[start:end].strip())
        if end == len(text):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]


def _extract_pdf_pages(path: Path) -> list[ExtractedPage]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required for PDF import. Run: pip install -r requirements.txt") from exc
    reader = PdfReader(str(path))
    pages: list[ExtractedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        text = _clean_text(page.extract_text() or "")
        if text:
            pages.append(ExtractedPage(page_number=index, text=text))
    return pages


def _extract_pdf(path: Path) -> str:
    return "\n\n".join(page.text for page in _extract_pdf_pages(path))


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for Word import. Run: pip install -r requirements.txt") from exc
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel import. Run: pip install -r requirements.txt") from exc
    workbook = load_workbook(path, data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                lines.append(f"Workbook: {path.name}. Sheet: {sheet.title}. Row {row_index}: " + " | ".join(values))
    return "\n".join(lines)


def _metadata_for_file(path: Path) -> dict[str, Any]:
    metadata_path = path.with_name(f"{path.name}.metadata.json")
    if not metadata_path.exists():
        return {}
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}
    return metadata if isinstance(metadata, dict) else {}


def _clean_pages(pages: list[ExtractedPage]) -> list[ExtractedPage]:
    return [ExtractedPage(page_number=page.page_number, text=_clean_text(page.text)) for page in pages if _clean_text(page.text)]


def _pdf_page_url(source_url: str | None, page_number: int | None) -> str | None:
    if not source_url or not page_number:
        return None
    base = source_url.split("#", 1)[0]
    return f"{base}#page={page_number}"


def _html_to_text(value: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\\1>", " ", value)
    value = re.sub(r"(?s)<[^>]+>", " ", value)
    return html.unescape(value)


def _title_from_html(value: str) -> str:
    match = re.search(r"(?is)<title[^>]*>(.*?)</title>", value)
    return _clean_text(match.group(1)) if match else ""


def _clean_text(value: str) -> str:
    value = re.sub(r"\r\n", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _hash_text(value: str) -> str:
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:24]
